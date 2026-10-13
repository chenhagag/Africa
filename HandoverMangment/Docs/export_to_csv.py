"""
Export apartment data from Excel status table to CSV for SharePoint import.
Source: מטריצת מעקב דירות - טבלת סטטוס sheet
Target: andreius3 list
"""
import openpyxl
import csv
import os

EXCEL_PATH = os.path.join(os.path.dirname(__file__), '- ליאור  מטריצת מעקב דירות - 22-06-2026 - דגם א.xlsx')
CSV_PATH = os.path.join(os.path.dirname(__file__), 'andreius3_import.csv')
SHEET_NAME = 'טבלת סטטוס'

# Excel column mapping (1-based index from sheet)
# C=Tower, D=Floor, E=AptNum, F=ShadowNum, G=Type, H=Rooms, I=Standard, K=Status, M=Handover, N=Finish
COL_MAP = {
    'Tower': 3,           # C - מגדל
    'Floor': 4,           # D - קומה
    'msrAppNum': 5,       # E - מספר דירה
    'ShadowNumber': 6,    # F - מספר צל
    'ApartmentType': 7,   # G - טיפוס
    'Rooms': 8,           # H - מספר חדרים
    'Standard': 9,        # I - סוג
    'msrAppStatus': 11,   # K - סטטוס
    'HandoverStatus': 13, # M - סטטוס מסירה
    'FinishStatus': 14,   # N - סטטוס גמר
}

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb[SHEET_NAME]

    rows_written = 0
    with open(CSV_PATH, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        # Header row - Title is required by SP
        headers = ['Title'] + list(COL_MAP.keys())
        writer.writerow(headers)

        # Data starts at row 3 (row 1 = merged header, row 2 = column headers)
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
            apt_num = row[COL_MAP['msrAppNum'] - 1].value
            if apt_num is None:
                continue

            values = [str(apt_num)]  # Title = apartment number
            for field, col_idx in COL_MAP.items():
                cell_val = row[col_idx - 1].value
                if cell_val is None:
                    values.append('')
                else:
                    values.append(str(cell_val).strip())

            writer.writerow(values)
            rows_written += 1

    wb.close()
    print(f'Exported {rows_written} apartments to {CSV_PATH}')

if __name__ == '__main__':
    main()
