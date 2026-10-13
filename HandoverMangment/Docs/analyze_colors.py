import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill

FILE = r"C:\SPFX\Africa\GIT-Africa\Africa\HandoverMangment\Docs\- ליאור  מטריצת מעקב דירות - 22-06-2026 - דגם א.xlsx"
SHEET = "מטריצת סטטוס עם הולד"

# Theme color index to name mapping (standard Office theme)
THEME_NAMES = {
    0: "Background1 (White)",
    1: "Text1 (Black)",
    2: "Background2 (Light Gray)",
    3: "Text2 (Dark Gray)",
    4: "Accent1 (Blue)",
    5: "Accent2 (Orange)",
    6: "Accent3 (Gray)",
    7: "Accent4 (Gold/Yellow)",
    8: "Accent5 (Blue-Gray)",
    9: "Accent6 (Green)",
}

def describe_color(color, label):
    """Describe a single color object safely."""
    if color is None:
        return f"{label}=None"
    try:
        theme = getattr(color, 'theme', None)
        tint = getattr(color, 'tint', None)
        rgb = getattr(color, 'rgb', None)
        indexed = getattr(color, 'indexed', None)
        ctype = getattr(color, 'type', None)

        if ctype == "rgb" and rgb and rgb != "00000000":
            return f"{label}=RGB({rgb})"
        elif theme is not None and isinstance(theme, int):
            tint_str = f", tint={tint:.4f}" if tint and tint != 0 else ""
            theme_name = THEME_NAMES.get(theme, f"theme{theme}")
            return f"{label}=Theme({theme}={theme_name}{tint_str})"
        elif indexed is not None and isinstance(indexed, int):
            return f"{label}=Indexed({indexed})"
        else:
            return f"{label}=({ctype},{rgb},{theme},{indexed})"
    except Exception as e:
        return f"{label}=ERROR({e})"


def describe_fill(cell):
    """Return a string describing the cell's fill color."""
    fill = cell.fill
    if fill is None:
        return "NO FILL"

    # Handle GradientFill
    from openpyxl.styles.fills import GradientFill
    if isinstance(fill, GradientFill) or (hasattr(fill, '__target') and isinstance(getattr(fill, '_StyleProxy__target', None), GradientFill)):
        try:
            return f"GRADIENT(type={getattr(fill, 'type', '?')}, stops={getattr(fill, 'stop', [])})"
        except:
            return "GRADIENT"

    try:
        pt = fill.patternType
    except AttributeError:
        return f"SPECIAL_FILL({type(fill).__name__})"

    if pt is None:
        return "NO FILL"

    parts = [f"pattern={pt}"]

    try:
        fg = fill.fgColor
    except:
        fg = None
    try:
        bg = fill.bgColor
    except:
        bg = None

    parts.append(describe_color(fg, "fgColor"))
    parts.append(describe_color(bg, "bgColor"))

    return " | ".join(parts)


print("Loading workbook...")
wb = openpyxl.load_workbook(FILE, data_only=True)
ws = wb[SHEET]

print(f"Sheet dimensions: {ws.dimensions}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
print(f"\nAnalyzing rows 191-227, columns A-BG (1-59)")
print("=" * 120)

# First pass: print column A values to identify row types
print("\n--- COLUMN A VALUES (row identifiers) ---")
for row in range(191, 228):
    val = ws.cell(row=row, column=1).value
    if val is not None:
        fill_info = describe_fill(ws.cell(row=row, column=1))
        print(f"  Row {row}: [{val}] | {fill_info}")

# Second pass: for each row, print all non-empty cells with colors
print("\n\n--- ALL NON-EMPTY CELLS WITH FILL COLORS ---")
for row in range(191, 228):
    row_label = ws.cell(row=row, column=1).value
    has_data = False
    for col in range(1, 60):  # A=1 through BG=59
        cell = ws.cell(row=row, column=col)
        val = cell.value
        if val is not None:
            if not has_data:
                print(f"\n{'='*100}")
                print(f"ROW {row} (Col A = {row_label})")
                print(f"{'='*100}")
                has_data = True
            col_letter = get_column_letter(col)
            fill_info = describe_fill(cell)
            print(f"  {col_letter}{row} = {repr(val)}  |  {fill_info}")

# Third pass: also check rows with no values but with fill colors (colored empty cells)
print("\n\n--- EMPTY CELLS WITH NON-DEFAULT FILL (rows 191-227) ---")
for row in range(191, 228):
    for col in range(1, 60):
        cell = ws.cell(row=row, column=col)
        if cell.value is None:
            fill = cell.fill
            if fill and fill.patternType and fill.patternType != "none":
                col_letter = get_column_letter(col)
                fill_info = describe_fill(cell)
                # Only print if it's not a trivial fill
                if "00000000" not in fill_info or "Theme" in fill_info:
                    print(f"  {col_letter}{row} (empty) | {fill_info}")

# Summary: unique fill colors found
print("\n\n--- UNIQUE FILL COLORS SUMMARY ---")
color_map = {}  # color_desc -> list of (row, col, value)
for row in range(191, 228):
    for col in range(1, 60):
        cell = ws.cell(row=row, column=col)
        fill_info = describe_fill(cell)
        if fill_info != "NO FILL":
            if fill_info not in color_map:
                color_map[fill_info] = []
            val = cell.value
            if val is not None:
                color_map[fill_info].append((row, get_column_letter(col), val))

for color_desc, cells in sorted(color_map.items()):
    print(f"\n  COLOR: {color_desc}")
    print(f"  Found in {len(cells)} non-empty cells. Examples:")
    for r, c, v in cells[:8]:
        print(f"    {c}{r} = {repr(v)}")

wb.close()
print("\nDone.")
