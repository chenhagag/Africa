import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles.fills import GradientFill, PatternFill

FILE = r"C:\SPFX\Africa\GIT-Africa\Africa\HandoverMangment\Docs\- ליאור  מטריצת מעקב דירות - 22-06-2026 - דגם א.xlsx"
SHEET = "מטריצת סטטוס עם הולד"

THEME_NAMES = {
    0: "White", 1: "Black", 2: "LightGray", 3: "DarkGray",
    4: "Blue", 5: "Orange", 6: "Gray", 7: "Gold",
    8: "BlueGray", 9: "Green",
}

def get_real_fill(cell):
    """Get the underlying fill object from a cell, handling StyleProxy."""
    fill = cell.fill
    target = getattr(fill, '_StyleProxy__target', fill)
    return target

def describe_fill(cell):
    fill = get_real_fill(cell)
    if fill is None:
        return "NO_FILL"
    if isinstance(fill, GradientFill):
        stops = []
        for s in (fill.stop or []):
            c = s.color
            if c and c.rgb:
                stops.append(f"RGB({c.rgb})")
            elif c and c.theme is not None:
                stops.append(f"Theme({c.theme})")
        return f"GRADIENT({','.join(stops)})"
    if isinstance(fill, PatternFill):
        if fill.patternType is None:
            return "NO_FILL"
        fg = fill.fgColor
        desc = f"pat={fill.patternType}"
        if fg:
            if fg.rgb and fg.rgb != "00000000" and fg.type == "rgb":
                desc += f" fg=RGB({fg.rgb})"
            elif fg.theme is not None and isinstance(fg.theme, int):
                t = f"tint={fg.tint:.2f}" if fg.tint else ""
                desc += f" fg=Theme({fg.theme}={THEME_NAMES.get(fg.theme,'?')}{','+t if t else ''})"
            elif fg.indexed is not None and isinstance(fg.indexed, int):
                desc += f" fg=Idx({fg.indexed})"
        return desc
    return f"UNKNOWN({type(fill).__name__})"

print("Loading workbook...")
wb = openpyxl.load_workbook(FILE, data_only=True)
ws = wb[SHEET]

print(f"Rows 191-227, Cols A-BG (1-59)\n")

# Collect unique colors with example cells
color_examples = {}  # color_desc -> [(row, col_letter, value), ...]

for row in range(191, 228):
    for col in range(1, 60):
        cell = ws.cell(row=row, column=col)
        val = cell.value
        fill_desc = describe_fill(cell)

        if val is not None:
            cl = get_column_letter(col)
            key = fill_desc
            if key not in color_examples:
                color_examples[key] = []
            color_examples[key].append((row, cl, str(val)[:60]))

# Print all non-empty cells grouped by row
print("=" * 130)
print("ALL NON-EMPTY CELLS (rows 191-227)")
print("=" * 130)
for row in range(191, 228):
    cells_in_row = []
    for col in range(1, 60):
        cell = ws.cell(row=row, column=col)
        if cell.value is not None:
            cl = get_column_letter(col)
            fill_desc = describe_fill(cell)
            cells_in_row.append(f"  {cl}{row} = {str(cell.value)[:50]}  |  {fill_desc}")
    if cells_in_row:
        print(f"\nROW {row}:")
        for c in cells_in_row:
            print(c)

# Summary
print("\n\n" + "=" * 130)
print("COLOR SUMMARY - unique fills with example values")
print("=" * 130)
for desc, examples in sorted(color_examples.items()):
    print(f"\n  [{desc}] - {len(examples)} cells")
    for r, c, v in examples[:10]:
        print(f"    {c}{r} = {v}")

wb.close()
print("\nDone.")
